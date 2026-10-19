"""单因子效果数据加载与查询。

数据来源：「因子效果.xlsx」（残酷+因子效果表）。
每行一个因子：因子名称、英文名、类别、分数、效果总结、详细说明。
详细说明为多行文本，以「·」「··」等前缀表示二、三级嵌套要点。
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .data import normalize_query

EXCEL_NAME = "因子效果.xlsx"
SHEET_NAME = "因子效果"
_SHEET_START = 2

CATEGORY_REGULAR = "常规"
CATEGORY_LIMITED = "限定"
CATEGORY_UNRELEASED = "未实装"
"""因子类别：常规（有分数）、节日限定与未实装（无分数）。"""


class FactorDataError(Exception):
    """因子效果数据加载或解析失败。"""


@dataclass(frozen=True)
class Factor:
    """一个突变因子的效果数据。"""

    name: str
    """因子中文名。"""

    en_name: str
    """因子英文名，未实装因子可能为空。"""

    category: str
    """类别：常规 / 限定 / 未实装。"""

    points: int | None
    """残酷+点数，限定与未实装因子为 None。"""

    summary: str
    """一句话效果总结。"""

    details: str
    """详细说明，多行文本；嵌套要点以「·」前缀表示。"""

    @property
    def points_text(self) -> str:
        """分数的展示文本，无分数时为空串。"""
        return f"{self.points}分" if self.points is not None else ""


def _parse_points(value: object) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(text)
    except ValueError as e:
        raise FactorDataError(f"无法解析因子分数：{value!r}") from e


def _parse_row(row: tuple[object, ...]) -> Factor | None:
    if row[0] is None or not str(row[0]).strip():
        return None
    name = str(row[0]).strip()
    try:
        return Factor(
            name=name,
            en_name="" if row[1] is None else str(row[1]).strip(),
            category="" if row[2] is None else str(row[2]).strip(),
            points=_parse_points(row[3]),
            summary="" if row[4] is None else str(row[4]).strip(),
            details="" if row[5] is None else str(row[5]).strip(),
        )
    except FactorDataError:
        raise
    except Exception as e:
        raise FactorDataError(f"解析因子行失败（{name}）：{e}") from e


class FactorBook:
    """全部因子效果数据，支持按名称查询。"""

    def __init__(self, factors: list[Factor]) -> None:
        if not factors:
            raise FactorDataError("因子数据为空")
        self._factors = factors
        self._by_name = {normalize_query(f.name): f for f in factors}

    @classmethod
    def load(cls, data_dir: Path) -> "FactorBook":
        """从数据目录读取因子效果 Excel。"""
        path = data_dir / EXCEL_NAME
        if not path.is_file():
            raise FactorDataError(f"未找到因子效果表：{path}")
        items: list[Factor] = []
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook[SHEET_NAME]
            for row in sheet.iter_rows(min_row=_SHEET_START, values_only=True):
                factor = _parse_row(row)
                if factor is not None:
                    items.append(factor)
        except FactorDataError:
            raise
        except Exception as e:
            raise FactorDataError(f"解析因子效果表失败：{e}") from e
        finally:
            if workbook is not None:
                workbook.close()
        return cls(items)

    @property
    def count(self) -> int:
        """数据中的因子个数。"""
        return len(self._factors)

    @property
    def factors(self) -> tuple[Factor, ...]:
        """全部因子，按表内顺序。"""
        return tuple(self._factors)

    def get(self, name: str) -> Factor | None:
        """按名称精确查询（忽略大小写与全/半角标点）。"""
        return self._by_name.get(normalize_query(name))

    def search(self, keyword: str) -> list[Factor]:
        """按名称或英文名模糊搜索，按表内顺序返回。

        关键字与名称均做全/半角标点归一化，任一宽度输入均可命中。
        """
        key = normalize_query(keyword)
        return [
            f
            for f in self._factors
            if key in normalize_query(f.name) or key in normalize_query(f.en_name)
        ]
