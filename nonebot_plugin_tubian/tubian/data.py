"""突变数据加载与选中逻辑。

数据来源：「突变列表数据.xlsx」（《星际争霸》每周突变数据表）。
每期突变自其「日期」（周一）起开放 7 天，按周轮换；
全部 184 期构成一轮完整循环，超出数据覆盖范围时按整轮折算。
"""

import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

EXCEL_NAME = "突变列表数据.xlsx"
ICON_DIR_NAME = "因子图标"
MAP_ICON_DIR_NAME = "地图图标"
WEEK_DAYS = 7
WEEKDAY_CN = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")

_LIST_SHEET = "突变列表"
_LIST_SHEET_START = 2
_CURRENT_SHEET = "本周突变"
_CURRENT_SHEET_START = 3


class MutationDataError(Exception):
    """突变数据加载或解析失败。"""


def local_today() -> date:
    """返回本地时区的今天。"""
    return datetime.now(tz=timezone.utc).astimezone().date()


def fmt_date(value: date) -> str:
    """格式化日期为 YYYY/MM/DD（周X）。"""
    day = f"{value.year}/{value.month:02d}/{value.day:02d}"
    return f"{day}（{WEEKDAY_CN[value.weekday()]}）"


def normalize_query(text: str) -> str:
    """查询文本归一化：去首尾空白、NFKC 折算全角为半角、统一小写。

    数据中的名称存在全/半角混用（如「我的机器人！」与「高举盾牌!」），
    归一化后按任一宽度输入均可命中。
    """
    return unicodedata.normalize("NFKC", text.strip()).lower()


@dataclass(frozen=True)
class Mutation:
    """一期每周突变。"""

    mid: int
    """轮换序号（1～184）。"""

    name: str
    """突变名称。"""

    map_name: str
    """地图名称。"""

    start_date: date
    """开放日期（当周周一）。"""

    factors: tuple[str, ...]
    """突变因子名称，1～3 个。"""

    number: str | None = None
    """原图右侧数字，缺失或为「/」时为 None。"""

    @property
    def end_date(self) -> date:
        """下一期开放日期，即本期限满之日。"""
        return self.start_date + timedelta(days=WEEK_DAYS)

    def remaining_days(self, today: date) -> int:
        """相对今天，本期剩余天数。"""
        return (self.end_date - today).days

    def contains(self, day: date) -> bool:
        """判断某天是否落在本期内。"""
        return self.start_date <= day < self.end_date


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("-", "/")
    parts = text.split("/")
    try:
        year, month, day = (int(part) for part in parts[:3])
        return date(year, month, day)
    except (ValueError, IndexError) as e:
        raise MutationDataError(f"无法解析日期：{value!r}") from e


def _parse_row(row: tuple[object, ...]) -> Mutation | None:
    if row[0] is None:
        return None
    try:
        mid = int(str(row[0]).strip())
    except ValueError as e:
        raise MutationDataError(f"无法解析序号：{row[0]!r}") from e
    factors = tuple(
        str(row[i]).strip()
        for i in (5, 7, 9)
        if row[i] is not None and str(row[i]).strip()
    )
    raw_number = "" if row[10] is None else str(row[10]).strip()
    number = raw_number if raw_number and raw_number not in {"/", "／", "-"} else None
    return Mutation(
        mid=mid,
        name=str(row[1]).strip(),
        map_name=str(row[2]).strip(),
        start_date=_parse_date(row[3]),
        factors=factors,
        number=number,
    )


class MutationBook:
    """全部突变数据，支持按时间、序号、名称选中。"""

    def __init__(
        self, mutations: list[Mutation], icon_dir: Path, map_icon_dir: Path
    ) -> None:
        by_id = {m.mid: m for m in mutations}
        if not by_id:
            raise MutationDataError("突变数据为空")
        self._by_id = by_id
        self._chrono = sorted(by_id.values(), key=lambda m: m.start_date)
        self.icon_dir = icon_dir
        self.map_icon_dir = map_icon_dir

    @classmethod
    def load(cls, data_dir: Path) -> "MutationBook":
        """从数据目录读取突变列表 Excel。"""
        path = data_dir / EXCEL_NAME
        if not path.is_file():
            raise MutationDataError(f"未找到突变数据表：{path}")
        items: dict[int, Mutation] = {}
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet_name, start in (
                (_LIST_SHEET, _LIST_SHEET_START),
                (_CURRENT_SHEET, _CURRENT_SHEET_START),
            ):
                sheet = workbook[sheet_name]
                for row in sheet.iter_rows(min_row=start, values_only=True):
                    mutation = _parse_row(row)
                    if mutation is not None:
                        # 「本周突变」表的数据更接近当前轮换，同序号时以此为准
                        items[mutation.mid] = mutation
        except MutationDataError:
            raise
        except Exception as e:
            raise MutationDataError(f"解析突变数据表失败：{e}") from e
        finally:
            if workbook is not None:
                workbook.close()
        return cls(
            list(items.values()),
            data_dir / ICON_DIR_NAME,
            data_dir / MAP_ICON_DIR_NAME,
        )

    @property
    def count(self) -> int:
        """数据中的突变期数。"""
        return len(self._by_id)

    def get(self, mid: int) -> Mutation | None:
        """按轮换序号查询。"""
        return self._by_id.get(mid)

    def search(self, keyword: str) -> list[Mutation]:
        """按突变名称或地图名称模糊搜索，按日期排序。

        关键字与名称均做全/半角标点归一化，任一宽度输入均可命中。
        """
        key = normalize_query(keyword)
        return [
            m
            for m in self._chrono
            if key in normalize_query(m.name) or key in normalize_query(m.map_name)
        ]

    def at(self, when: date) -> Mutation:
        """选中 when 所在周（周一开放、为期 7 天）的突变。

        超出数据覆盖范围时按完整轮换周期折算（循环 184 周）。
        """
        for m in self._chrono:
            if m.contains(when):
                return m
        anchor = self._chrono[0].start_date
        idx = ((when - anchor).days // WEEK_DAYS) % len(self._chrono)
        return self._chrono[idx]

    def icon_path(self, factor: str) -> Path | None:
        """返回因子图标路径，不存在则返回 None。"""
        path = self.icon_dir / f"因子_{factor}.png"
        return path if path.is_file() else None

    def map_icon_path(self, map_name: str) -> Path | None:
        """返回地图图片路径，不存在则返回 None。"""
        path = self.map_icon_dir / f"{map_name}.png"
        return path if path.is_file() else None
